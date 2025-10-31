# ==============================
# INÍCIO: db_handler.py COMPLETO FINAL CORRIGIDO
# ==============================

import sqlite3
import logging
import os
import shutil
from typing import Dict, List, Any, Optional, Tuple
from datetime import datetime, date
import tempfile

logger = logging.getLogger(__name__)

def get_db_connection(db_path: str) -> sqlite3.Connection:
    """Cria conexão com o banco SQLite com configurações otimizadas"""
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        # Configurações para melhor performance
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
        conn.execute("PRAGMA foreign_keys=ON")
        return conn
    except sqlite3.Error as e:
        logger.error(f"Erro ao conectar com banco {db_path}: {str(e)}")
        raise

# ==============================
# FUNÇÃO VERIFICAR_LOCALIDADE_EXISTE - ADICIONADA PARA CORRIGIR ERRO
# ==============================

def verificar_localidade_existe(db_path: str, localidade: str) -> Dict[str, Any]:
    """
    Verifica se uma localidade existe no banco e retorna informações
    Retorna: Dict com informações sobre a localidade
    """
    conn = None
    try:
        if not localidade or not localidade.strip():
            return {
                'existe': False,
                'mensagem': 'Localidade não informada',
                'quantidade_bens': 0,
                'ultima_atualizacao': None
            }
        
        localidade_limpa = localidade.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se existem bens nessa localidade
        cursor.execute('''
            SELECT 
                COUNT(*) as quantidade,
                MAX(ultima_atualizacao) as ultima_atualizacao
            FROM bens 
            WHERE localizacao = ?
        ''', (localidade_limpa,))
        
        resultado = cursor.fetchone()
        quantidade = resultado['quantidade'] if resultado else 0
        ultima_atualizacao = resultado['ultima_atualizacao'] if resultado else None
        
        if quantidade > 0:
            return {
                'existe': True,
                'mensagem': f'Localidade encontrada com {quantidade} bem(ns)',
                'quantidade_bens': quantidade,
                'ultima_atualizacao': ultima_atualizacao,
                'localidade': localidade_limpa
            }
        else:
            return {
                'existe': False,
                'mensagem': 'Nenhum bem encontrado para esta localidade',
                'quantidade_bens': 0,
                'ultima_atualizacao': None,
                'localidade': localidade_limpa
            }
            
    except Exception as e:
        logger.error(f"Erro ao verificar localidade '{localidade}': {str(e)}")
        return {
            'existe': False,
            'mensagem': f'Erro ao verificar localidade: {str(e)}',
            'quantidade_bens': 0,
            'ultima_atualizacao': None
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO VERIFICAR_NUMERO_EXISTE - ADICIONADA PARA CORRIGIR ERRO DE IMPORTAÇÃO
# ==============================

def verificar_numero_existe(db_path: str, numero_bem: str) -> Dict[str, Any]:
    """
    Verifica se um número de bem já existe no banco
    Retorna: Dict com informação se o número existe
    """
    conn = None
    try:
        if not numero_bem or not numero_bem.strip():
            return {
                'existe': False,
                'mensagem': 'Número do bem não informado',
                'numero_bem': numero_bem
            }
        
        numero_limpo = numero_bem.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT numero, nome, situacao
            FROM bens 
            WHERE numero = ?
        ''', (numero_limpo,))
        
        bem = cursor.fetchone()
        
        if bem:
            bem_dict = dict(bem)
            return {
                'existe': True,
                'mensagem': f'Bem {numero_limpo} já existe',
                'numero_bem': numero_limpo,
                'bem': bem_dict
            }
        else:
            return {
                'existe': False,
                'mensagem': f'Bem {numero_limpo} não encontrado',
                'numero_bem': numero_limpo,
                'bem': None
            }
            
    except Exception as e:
        logger.error(f"Erro ao verificar número {numero_bem}: {str(e)}")
        return {
            'existe': False,
            'mensagem': f'Erro ao verificar número: {str(e)}',
            'numero_bem': numero_bem,
            'bem': None
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO OBTER_BEM_POR_ID - ADICIONADA PARA CORRIGIR ERRO DE IMPORTAÇÃO
# ==============================

def obter_bem_por_id(db_path: str, bem_id: int) -> Dict[str, Any]:
    """
    Obtém um bem específico pelo ID
    Retorna: Dict com informações do bem ou erro
    """
    conn = None
    try:
        if not bem_id or bem_id <= 0:
            return {
                'sucesso': False,
                'mensagem': 'ID do bem inválido',
                'bem': None
            }
        
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT 
                id, numero, nome, situacao, localizacao, responsavel,
                data_ultima_vistoria, data_vistoria_atual, auditor,
                observacoes, data_criacao, data_localizacao
            FROM bens 
            WHERE id = ?
        ''', (bem_id,))
        
        bem = cursor.fetchone()
        
        if bem:
            # Converter para dict
            bem_dict = dict(bem)
            
            # Converter datas para string para serialização
            for key, value in bem_dict.items():
                if isinstance(value, (datetime, date)):
                    bem_dict[key] = value.isoformat()
            
            return {
                'sucesso': True,
                'mensagem': 'Bem encontrado',
                'bem': bem_dict
            }
        else:
            return {
                'sucesso': False,
                'mensagem': f'Bem com ID {bem_id} não encontrado',
                'bem': None
            }
            
    except sqlite3.Error as e:
        logger.error(f"Erro ao obter bem por ID {bem_id}: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro no banco de dados: {str(e)}',
            'bem': None
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao obter bem por ID {bem_id}: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}',
            'bem': None
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO CRIAR_TABELA_ATUALIZADA - ADICIONADA PARA CORRIGIR ERRO DE IMPORTAÇÃO
# ==============================

def criar_tabela_atualizada(db_path: str) -> Dict[str, Any]:
    """
    Cria a tabela bens com estrutura atualizada e migra dados se necessário
    Retorna: Dict com resultado da operação
    """
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se a tabela existe
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='bens'")
        tabela_existe = cursor.fetchone() is not None
        
        if not tabela_existe:
            # Criar tabela do zero
            sucesso = criar_tabela_se_nao_existir(db_path)
            return {
                'sucesso': sucesso,
                'mensagem': 'Tabela bens criada com sucesso' if sucesso else 'Erro ao criar tabela',
                'tabela_criada': sucesso,
                'migracao_realizada': False
            }
        else:
            # Tabela já existe, verificar e adicionar colunas faltantes
            colunas_necessarias = [
                ('responsavel', 'TEXT'),
                ('data_ultima_vistoria', 'DATE'),
                ('data_vistoria_atual', 'DATE'),
                ('auditor', 'TEXT'),
                ('ultima_atualizacao', 'DATETIME'),
                ('data_localizacao', 'DATETIME')
            ]
            
            cursor.execute("PRAGMA table_info(bens)")
            colunas_existentes = [col[1] for col in cursor.fetchall()]
            
            colunas_adicionadas = []
            for coluna, tipo in colunas_necessarias:
                if coluna not in colunas_existentes:
                    try:
                        cursor.execute(f"ALTER TABLE bens ADD COLUMN {coluna} {tipo}")
                        colunas_adicionadas.append(coluna)
                        logger.info(f"Coluna {coluna} adicionada à tabela bens")
                    except sqlite3.OperationalError as e:
                        if "duplicate column name" not in str(e).lower():
                            logger.warning(f"Erro ao adicionar coluna {coluna}: {e}")
            
            # Criar índices para performance se não existirem
            try:
                cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_numero ON bens(numero)")
                cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_situacao ON bens(situacao)")
                cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_localizacao ON bens(localizacao)")
                cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_nome ON bens(nome)")
            except sqlite3.Error as e:
                logger.warning(f"Erro ao criar índices: {e}")
            
            conn.commit()
            
            mensagem = "Tabela bens verificada com sucesso"
            if colunas_adicionadas:
                mensagem += f". Colunas adicionadas: {', '.join(colunas_adicionadas)}"
            
            return {
                'sucesso': True,
                'mensagem': mensagem,
                'tabela_criada': False,  # Já existia
                'migracao_realizada': len(colunas_adicionadas) > 0,
                'colunas_adicionadas': colunas_adicionadas
            }
            
    except Exception as e:
        logger.error(f"Erro ao criar tabela atualizada: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro ao criar tabela: {str(e)}',
            'tabela_criada': False,
            'migracao_realizada': False
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO OBTER_BENS_POR_LOCALIDADE - CORRIGIDA
# ==============================

def obter_bens_por_localidade(db_path: str, localidade: str, pagina: int = 1, por_pagina: int = 50) -> Dict[str, Any]:
    """
    Obtém bens por localidade específica com paginação
    Retorna: Dict com lista de bens da localidade
    """
    conn = None
    try:
        if not localidade or not localidade.strip():
            return {
                'sucesso': False,
                'mensagem': 'Localidade não informada',
                'bens': [],
                'total': 0
            }
        
        localidade_limpa = localidade.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Query base para contar total
        count_query = '''
            SELECT COUNT(*) 
            FROM bens 
            WHERE localizacao = ?
        '''
        cursor.execute(count_query, (localidade_limpa,))
        total = cursor.fetchone()[0]
        
        if total == 0:
            return {
                'sucesso': True,
                'mensagem': f'Nenhum bem encontrado para a localidade: {localidade_limpa}',
                'bens': [],
                'total': 0,
                'localidade': localidade_limpa
            }
        
        # Query para obter bens com paginação
        offset = (pagina - 1) * por_pagina
        query = '''
            SELECT 
                id, numero, nome, situacao, localizacao, responsavel,
                data_ultima_vistoria, data_vistoria_atual, auditor,
                observacoes, data_criacao, data_localizacao
            FROM bens 
            WHERE localizacao = ?
            ORDER BY numero
            LIMIT ? OFFSET ?
        '''
        
        cursor.execute(query, (localidade_limpa, por_pagina, offset))
        bens = [dict(row) for row in cursor.fetchall()]
        
        # Converter datas para string
        for bem in bens:
            for key, value in bem.items():
                if isinstance(value, (datetime, date)):
                    bem[key] = value.isoformat()
        
        return {
            'sucesso': True,
            'bens': bens,
            'total': total,
            'pagina': pagina,
            'por_pagina': por_pagina,
            'total_paginas': (total + por_pagina - 1) // por_pagina if por_pagina > 0 else 1,
            'localidade': localidade_limpa
        }
        
    except Exception as e:
        logger.error(f"Erro ao obter bens por localidade '{localidade}': {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro ao buscar bens: {str(e)}',
            'bens': [],
            'total': 0
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO OBTER_TODOS_BENS_POR_LOCALIDADE - ADICIONADA PARA CORRIGIR O ERRO
# ==============================

def obter_todos_bens_por_localidade(db_path: str, localidade: str) -> Dict[str, Any]:
    """
    Obtém TODOS os bens de uma localidade específica (sem paginação)
    Esta função é necessária para corrigir o erro de importação
    """
    conn = None
    try:
        if not localidade or not localidade.strip():
            return {
                'sucesso': False,
                'mensagem': 'Localidade não informada',
                'bens': [],
                'total': 0
            }
        
        localidade_limpa = localidade.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Query para obter todos os bens da localidade
        query = '''
            SELECT 
                id, numero, nome, situacao, localizacao, responsavel,
                data_ultima_vistoria, data_vistoria_atual, auditor,
                observacoes, data_criacao, data_localizacao
            FROM bens 
            WHERE localizacao = ?
            ORDER BY numero
        '''
        
        cursor.execute(query, (localidade_limpa,))
        bens = [dict(row) for row in cursor.fetchall()]
        
        # Converter datas para string
        for bem in bens:
            for key, value in bem.items():
                if isinstance(value, (datetime, date)):
                    bem[key] = value.isoformat()
        
        return {
            'sucesso': True,
            'bens': bens,
            'total': len(bens),
            'localidade': localidade_limpa
        }
        
    except Exception as e:
        logger.error(f"Erro ao obter todos bens por localidade '{localidade}': {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro ao buscar bens: {str(e)}',
            'bens': [],
            'total': 0
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO OBTER_LOCALIDADES - ADICIONADA
# ==============================

def obter_localidades(db_path: str) -> Dict[str, Any]:
    """
    Obtém lista de localidades únicas (alias para obter_localizacoes)
    Retorna: Dict com lista de localidades
    """
    # Reutiliza a função obter_localizacoes existente
    localizacoes = obter_localizacoes(db_path)
    return {
        'sucesso': True,
        'localidades': localizacoes,
        'total': len(localizacoes)
    }

# ==============================
# FUNÇÃO OBTER_BENS_PAGINADOS - ADICIONADA
# ==============================

def obter_bens_paginados(db_path: str, pagina: int = 1, por_pagina: int = 50, filtros: Dict[str, Any] = None) -> Dict[str, Any]:
    """
    Obtém bens com paginação e filtros (alias para buscar_bens)
    Retorna: Dict com lista de bens e informações de paginação
    """
    # Reutiliza a função buscar_bens existente
    return buscar_bens(db_path, filtros, pagina, por_pagina)

# ==============================
# FUNÇÃO CONTAR_BENS - ADICIONADA
# ==============================

def contar_bens(db_path: str, filtros: Dict[str, Any] = None) -> Dict[str, Any]:
    """
    Conta o total de bens com filtros opcionais
    Retorna: Dict com contagens detalhadas
    """
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Query base para contar todos os bens
        query_base = "SELECT COUNT(*) FROM bens WHERE 1=1"
        params = []
        
        # Aplicar filtros se fornecidos
        if filtros:
            if filtros.get('numero'):
                query_base += " AND numero LIKE ?"
                params.append(f"%{filtros['numero']}%")
            
            if filtros.get('nome'):
                query_base += " AND nome LIKE ?"
                params.append(f"%{filtros['nome']}%")
            
            if filtros.get('situacao'):
                query_base += " AND situacao = ?"
                params.append(filtros['situacao'])
            
            if filtros.get('localizacao'):
                query_base += " AND localizacao LIKE ?"
                params.append(f"%{filtros['localizacao']}%")
        
        # Contar total
        cursor.execute(query_base, params)
        total = cursor.fetchone()[0]
        
        # Contar por situação
        situacoes_query = """
            SELECT situacao, COUNT(*) as quantidade 
            FROM bens 
            WHERE 1=1
        """
        situacoes_params = []
        
        if filtros:
            if filtros.get('numero'):
                situacoes_query += " AND numero LIKE ?"
                situacoes_params.append(f"%{filtros['numero']}%")
            
            if filtros.get('nome'):
                situacoes_query += " AND nome LIKE ?"
                situacoes_params.append(f"%{filtros['nome']}%")
            
            if filtros.get('localizacao'):
                situacoes_query += " AND localizacao LIKE ?"
                situacoes_params.append(f"%{filtros['localizacao']}%")
        
        situacoes_query += " GROUP BY situacao ORDER BY quantidade DESC"
        
        cursor.execute(situacoes_query, situacoes_params)
        situacoes = {row['situacao']: row['quantidade'] for row in cursor.fetchall()}
        
        # Contar por localização (top 10)
        localizacoes_query = """
            SELECT localizacao, COUNT(*) as quantidade 
            FROM bens 
            WHERE localizacao IS NOT NULL AND localizacao != ''
        """
        localizacoes_params = []
        
        if filtros:
            if filtros.get('numero'):
                localizacoes_query += " AND numero LIKE ?"
                localizacoes_params.append(f"%{filtros['numero']}%")
            
            if filtros.get('nome'):
                localizacoes_query += " AND nome LIKE ?"
                localizacoes_params.append(f"%{filtros['nome']}%")
            
            if filtros.get('situacao'):
                localizacoes_query += " AND situacao = ?"
                localizacoes_params.append(filtros['situacao'])
        
        localizacoes_query += " GROUP BY localizacao ORDER BY quantidade DESC LIMIT 10"
        
        cursor.execute(localizacoes_query, localizacoes_params)
        localizacoes = {row['localizacao']: row['quantidade'] for row in cursor.fetchall()}
        
        return {
            'sucesso': True,
            'total': total,
            'situacoes': situacoes,
            'localizacoes': localizacoes,
            'filtros_aplicados': bool(filtros)
        }
        
    except Exception as e:
        logger.error(f"Erro ao contar bens: {str(e)}")
        return {
            'sucesso': False,
            'total': 0,
            'situacoes': {},
            'localizacoes': {},
            'mensagem': f'Erro ao contar bens: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO BUSCAR_BENS_POR_NOME - ADICIONADA
# ==============================

def buscar_bens_por_nome(db_path: str, termo: str, limite: int = 50) -> Dict[str, Any]:
    """
    Busca bens por nome (para autocomplete e buscas rápidas)
    Retorna: Dict com lista de bens encontrados
    """
    conn = None
    try:
        if not termo or not termo.strip():
            return {
                'sucesso': True,
                'bens': [],
                'total': 0,
                'termo_busca': termo
            }
        
        termo_limpo = termo.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Buscar bens que contenham o termo no nome
        cursor.execute('''
            SELECT numero, nome, situacao, localizacao
            FROM bens 
            WHERE nome LIKE ?
            ORDER BY 
                CASE 
                    WHEN nome LIKE ? THEN 1  -- Prioridade para começo do nome
                    ELSE 2
                END,
                nome
            LIMIT ?
        ''', (f'%{termo_limpo}%', f'{termo_limpo}%', limite))
        
        bens = [dict(row) for row in cursor.fetchall()]
        
        return {
            'sucesso': True,
            'bens': bens,
            'total': len(bens),
            'termo_busca': termo_limpo
        }
        
    except Exception as e:
        logger.error(f"Erro ao buscar bens por nome '{termo}': {str(e)}")
        return {
            'sucesso': False,
            'bens': [],
            'total': 0,
            'mensagem': f'Erro ao buscar bens: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO CRIAR_NOVO_BEM - CONTINUAÇÃO
# ==============================

def criar_novo_bem(db_path: str, dados: Dict[str, Any]) -> Dict[str, Any]:
    """
    Cria um novo bem no banco de dados
    Retorna: Dict com resultado da operação
    """
    conn = None
    try:
        # Validar dados obrigatórios
        if not dados.get('numero') or not dados['numero'].strip():
            return {
                'sucesso': False,
                'mensagem': 'Número do bem é obrigatório'
            }
        
        if not dados.get('nome') or not dados['nome'].strip():
            return {
                'sucesso': False,
                'mensagem': 'Nome do bem é obrigatório'
            }
        
        numero_limpo = dados['numero'].strip()
        nome_limpo = dados['nome'].strip()
        
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se bem já existe
        cursor.execute("SELECT id FROM bens WHERE numero = ?", (numero_limpo,))
        if cursor.fetchone():
            return {
                'sucesso': False,
                'mensagem': f'Bem com número {numero_limpo} já existe'
            }
        
        # Preparar dados para inserção
        campos = ['numero', 'nome']
        valores = [numero_limpo, nome_limpo]
        
        # Campos opcionais
        campos_opcionais = [
            'situacao', 'localizacao', 'responsavel', 
            'data_ultima_vistoria', 'data_vistoria_atual', 
            'auditor', 'observacoes'
        ]
        
        for campo in campos_opcionais:
            if campo in dados and dados[campo] is not None:
                campos.append(campo)
                valores.append(dados[campo])
        
        # Inserir novo bem
        placeholders = ', '.join(['?' for _ in campos])
        campos_str = ', '.join(campos)
        
        query = f"INSERT INTO bens ({campos_str}) VALUES ({placeholders})"
        cursor.execute(query, valores)
        
        conn.commit()
        
        mensagem = f"Bem {numero_limpo} - {nome_limpo} criado com sucesso"
        logger.info(f"NOVO BEM CRIADO: {mensagem}")
        
        return {
            'sucesso': True,
            'mensagem': mensagem,
            'numero_bem': numero_limpo,
            'id': cursor.lastrowid
        }
        
    except sqlite3.Error as e:
        logger.error(f"Erro ao criar bem {dados.get('numero')}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro no banco de dados: {str(e)}'
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao criar bem {dados.get('numero')}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO EXCLUIR_BEM - CONTINUAÇÃO
# ==============================

def excluir_bem(db_path: str, numero_bem: str) -> Dict[str, Any]:
    """
    Exclui um bem do banco de dados pelo número
    Retorna: Dict com resultado da operação
    """
    conn = None
    try:
        if not numero_bem or not numero_bem.strip():
            return {
                'sucesso': False,
                'mensagem': 'Número do bem não informado'
            }
        
        numero_limpo = numero_bem.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se bem existe antes de excluir
        cursor.execute("SELECT id, numero, nome FROM bens WHERE numero = ?", (numero_limpo,))
        bem = cursor.fetchone()
        
        if not bem:
            return {
                'sucesso': False,
                'mensagem': f'Bem {numero_limpo} não encontrado'
            }
        
        # Registrar informações do bem antes de excluir (para log)
        bem_info = f"{bem['numero']} - {bem['nome']}"
        
        # Excluir o bem
        cursor.execute("DELETE FROM bens WHERE numero = ?", (numero_limpo,))
        
        conn.commit()
        
        mensagem = f"Bem {bem_info} excluído com sucesso"
        logger.warning(f"BEM EXCLUÍDO: {mensagem}")
        
        return {
            'sucesso': True,
            'mensagem': mensagem,
            'numero_bem': numero_limpo
        }
        
    except sqlite3.Error as e:
        logger.error(f"Erro ao excluir bem {numero_bem}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro no banco de dados: {str(e)}'
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao excluir bem {numero_bem}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO OBTER_BEM_POR_NUMERO - CONTINUAÇÃO
# ==============================

def obter_bem_por_numero(db_path: str, numero_bem: str) -> Dict[str, Any]:
    """
    Obtém um bem específico pelo número (alias para verificar_bem)
    Retorna: Dict com informações do bem ou erro
    """
    # Reutiliza a função verificar_bem existente
    return verificar_bem(db_path, numero_bem)

# ==============================
# FUNÇÃO BUSCAR_LOCALIZACAO_EXISTENTE - CONTINUAÇÃO
# ==============================

def buscar_localizacao_existente(db_path: str, termo: str = None) -> Dict[str, Any]:
    """
    Busca localizações existentes no banco com sugestões
    Retorna: Dict com lista de localizações
    """
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        if termo and termo.strip():
            # Buscar localizações que contenham o termo
            cursor.execute('''
                SELECT DISTINCT localizacao, COUNT(*) as quantidade
                FROM bens 
                WHERE localizacao IS NOT NULL 
                AND localizacao != ''
                AND localizacao LIKE ?
                GROUP BY localizacao
                ORDER BY quantidade DESC, localizacao
                LIMIT 20
            ''', (f'%{termo.strip()}%',))
        else:
            # Buscar todas as localizações (mais frequentes primeiro)
            cursor.execute('''
                SELECT DISTINCT localizacao, COUNT(*) as quantidade
                FROM bens 
                WHERE localizacao IS NOT NULL 
                AND localizacao != ''
                GROUP BY localizacao
                ORDER BY quantidade DESC, localizacao
                LIMIT 50
            ''')
        
        resultados = cursor.fetchall()
        localizacoes = [{
            'nome': row['localizacao'],
            'quantidade': row['quantidade']
        } for row in resultados]
        
        return {
            'sucesso': True,
            'localizacoes': localizacoes,
            'total': len(localizacoes),
            'termo_busca': termo
        }
        
    except Exception as e:
        logger.error(f"Erro ao buscar localizações: {str(e)}")
        return {
            'sucesso': False,
            'localizacoes': [],
            'total': 0,
            'mensagem': f'Erro ao buscar localizações: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO GERAR_PLANILHAS_LOCALIZACAO - CONTINUAÇÃO
# ==============================

def gerar_planilhas_localizacao(db_path: str, output_dir: str = "relatorios") -> Dict[str, Any]:
    """
    Gera planilhas separadas por localização
    Retorna: Dict com resultado da operação
    """
    try:
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        # Criar diretório de saída se não existir
        os.makedirs(output_dir, exist_ok=True)
        
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Obter todas as localizações únicas
        cursor.execute('''
            SELECT DISTINCT localizacao 
            FROM bens 
            WHERE localizacao IS NOT NULL AND localizacao != ''
            ORDER BY localizacao
        ''')
        
        localizacoes = [row[0] for row in cursor.fetchall()]
        
        if not localizacoes:
            return {
                'sucesso': False,
                'mensagem': 'Nenhuma localização encontrada no banco de dados',
                'arquivos_gerados': []
            }
        
        arquivos_gerados = []
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        # Gerar planilha para cada localização
        for localizacao in localizacoes:
            try:
                # Buscar bens da localização
                cursor.execute('''
                    SELECT 
                        numero, nome, situacao, localizacao, responsavel,
                        data_ultima_vistoria, data_vistoria_atual, auditor,
                        observacoes, data_criacao, data_localizacao
                    FROM bens 
                    WHERE localizacao = ?
                    ORDER BY numero
                ''', (localizacao,))
                
                bens = [dict(row) for row in cursor.fetchall()]
                
                if not bens:
                    continue
                
                # Criar DataFrame
                df = pd.DataFrame(bens)
                
                # Nome do arquivo (remover caracteres inválidos)
                nome_arquivo = f"localizacao_{localizacao}_{timestamp}.xlsx"
                nome_arquivo = "".join(c for c in nome_arquivo if c.isalnum() or c in (' ', '-', '_', '.')).rstrip()
                caminho_arquivo = os.path.join(output_dir, nome_arquivo)
                
                # Exportar para Excel
                with pd.ExcelWriter(caminho_arquivo, engine='openpyxl') as writer:
                    df.to_excel(writer, sheet_name=localizacao[:31], index=False)
                    
                    # Formatar planilha
                    workbook = writer.book
                    worksheet = writer.sheets[localizacao[:31]]
                    
                    # Formatar cabeçalho
                    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                    header_font = Font(color="FFFFFF", bold=True)
                    
                    for col in range(1, len(df.columns) + 1):
                        cell = worksheet.cell(row=1, column=col)
                        cell.fill = header_fill
                        cell.font = header_font
                    
                    # Ajustar largura das colunas
                    for idx, col in enumerate(df.columns, 1):
                        max_len = max(df[col].astype(str).str.len().max(), len(str(col))) + 2
                        worksheet.column_dimensions[chr(64 + idx)].width = min(max_len, 50)
                
                arquivos_gerados.append({
                    'localizacao': localizacao,
                    'arquivo': nome_arquivo,
                    'caminho': caminho_arquivo,
                    'quantidade': len(bens)
                })
                
                logger.info(f"Planilha gerada para {localizacao}: {len(bens)} bens")
                
            except Exception as e:
                logger.error(f"Erro ao gerar planilha para {localizacao}: {str(e)}")
                continue
        
        conn.close()
        
        if not arquivos_gerados:
            return {
                'sucesso': False,
                'mensagem': 'Nenhuma planilha pôde ser gerada',
                'arquivos_gerados': []
            }
        
        # Gerar planilha consolidada com todas as localizações
        try:
            caminho_consolidado = os.path.join(output_dir, f"consolidado_localizacoes_{timestamp}.xlsx")
            
            with pd.ExcelWriter(caminho_consolidado, engine='openpyxl') as writer:
                for localizacao in localizacoes:
                    cursor.execute('''
                        SELECT numero, nome, situacao, localizacao, responsavel
                        FROM bens WHERE localizacao = ? ORDER BY numero
                    ''', (localizacao,))
                    
                    bens = [dict(row) for row in cursor.fetchall()]
                    if bens:
                        df_local = pd.DataFrame(bens)
                        sheet_name = localizacao[:31]  # Limitar a 31 caracteres
                        df_local.to_excel(writer, sheet_name=sheet_name, index=False)
            
            arquivos_gerados.append({
                'localizacao': 'CONSOLIDADO',
                'arquivo': f"consolidado_localizacoes_{timestamp}.xlsx",
                'caminho': caminho_consolidado,
                'quantidade': len(localizacoes)
            })
        except Exception as e:
            logger.error(f"Erro ao gerar planilha consolidada: {str(e)}")
        
        return {
            'sucesso': True,
            'mensagem': f'Geradas {len(arquivos_gerados)} planilhas de localização',
            'arquivos_gerados': arquivos_gerados,
            'total_localizacoes': len(localizacoes)
        }
        
    except ImportError as e:
        logger.error(f"Bibliotecas necessárias não instaladas: {e}")
        return {
            'sucesso': False,
            'mensagem': 'Bibliotecas pandas/openpyxl não instaladas',
            'arquivos_gerados': []
        }
    except Exception as e:
        logger.error(f"Erro ao gerar planilhas de localização: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro ao gerar planilhas: {str(e)}',
            'arquivos_gerados': []
        }

# ==============================
# FUNÇÃO VERIFICAR_BEM - CONTINUAÇÃO
# ==============================

def verificar_bem(db_path: str, numero_bem: str) -> Dict[str, Any]:
    """
    Verifica se um bem existe no banco e retorna seus dados
    Retorna: Dict com informações do bem ou erro
    """
    conn = None
    try:
        if not numero_bem or not numero_bem.strip():
            return {
                'sucesso': False,
                'mensagem': 'Número do bem não informado',
                'bem': None
            }
        
        numero_limpo = numero_bem.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT 
                id, numero, nome, situacao, localizacao, responsavel,
                data_ultima_vistoria, data_vistoria_atual, auditor,
                observacoes, data_criacao, data_localizacao
            FROM bens 
            WHERE numero = ?
        ''', (numero_limpo,))
        
        bem = cursor.fetchone()
        
        if bem:
            # Converter para dict
            bem_dict = dict(bem)
            
            # Converter datas para string para serialização
            for key, value in bem_dict.items():
                if isinstance(value, (datetime, date)):
                    bem_dict[key] = value.isoformat()
            
            return {
                'sucesso': True,
                'mensagem': 'Bem encontrado',
                'bem': bem_dict
            }
        else:
            return {
                'sucesso': False,
                'mensagem': f'Bem {numero_limpo} não encontrado',
                'bem': None
            }
            
    except sqlite3.Error as e:
        logger.error(f"Erro ao verificar bem {numero_bem}: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro no banco de dados: {str(e)}',
            'bem': None
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao verificar bem {numero_bem}: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}',
            'bem': None
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÃO MARCAR_BEM_LOCALIZADO - CONTINUAÇÃO
# ==============================

def marcar_bem_localizado(db_path: str, numero_bem: str, localizacao: str = None, observacoes: str = None) -> Dict[str, Any]:
    """
    Marca um bem como localizado e atualiza sua localização
    Retorna: Dict com resultado da operação
    """
    conn = None
    try:
        if not numero_bem or not numero_bem.strip():
            return {
                'sucesso': False,
                'mensagem': 'Número do bem não informado'
            }
        
        numero_limpo = numero_bem.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se bem existe
        cursor.execute("SELECT id, localizacao FROM bens WHERE numero = ?", (numero_limpo,))
        bem = cursor.fetchone()
        
        if not bem:
            return {
                'sucesso': False,
                'mensagem': f'Bem {numero_limpo} não encontrado'
            }
        
        # Preparar dados para atualização
        campos_atualizar = []
        parametros = []
        
        # Sempre marcar como "Localizado"
        campos_atualizar.append("situacao = ?")
        parametros.append("Localizado")
        
        # Atualizar localização se fornecida
        if localizacao and localizacao.strip():
            campos_atualizar.append("localizacao = ?")
            parametros.append(localizacao.strip())
        
        # Atualizar observações se fornecidas
        if observacoes is not None:
            campos_atualizar.append("observacoes = ?")
            parametros.append(observacoes.strip() if observacoes.strip() else "")
        
        # Sempre atualizar data_localizacao
        campos_atualizar.append("data_localizacao = CURRENT_TIMESTAMP")
        campos_atualizar.append("ultima_atualizacao = CURRENT_TIMESTAMP")
        
        parametros.append(numero_limpo)
        
        # Executar atualização
        query = f"UPDATE bens SET {', '.join(campos_atualizar)} WHERE numero = ?"
        cursor.execute(query, parametros)
        
        conn.commit()
        
        mensagem = f"Bem {numero_limpo} marcado como localizado"
        if localizacao and localizacao.strip():
            mensagem += f" na localização: {localizacao.strip()}"
        
        logger.info(mensagem)
        
        return {
            'sucesso': True,
            'mensagem': mensagem,
            'numero_bem': numero_limpo,
            'localizacao': localizacao.strip() if localizacao else None
        }
        
    except sqlite3.Error as e:
        logger.error(f"Erro ao marcar bem {numero_bem} como localizado: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro no banco de dados: {str(e)}'
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao marcar bem {numero_bem}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FUNÇÕES EXISTENTES (MANTIDAS E CORRIGIDAS)
# ==============================

def criar_tabela_se_nao_existir(db_path: str) -> bool:
    """Cria a tabela bens se não existir"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS bens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                numero TEXT UNIQUE NOT NULL,
                nome TEXT NOT NULL,
                situacao TEXT DEFAULT 'Pendente',
                localizacao TEXT,
                responsavel TEXT,
                data_ultima_vistoria DATE,
                data_vistoria_atual DATE,
                auditor TEXT,
                observacoes TEXT,
                data_criacao DATETIME DEFAULT CURRENT_TIMESTAMP,
                data_localizacao DATETIME,
                ultima_atualizacao DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # Verificar e adicionar colunas faltantes
        colunas_necessarias = [
            ('responsavel', 'TEXT'),
            ('data_ultima_vistoria', 'DATE'),
            ('data_vistoria_atual', 'DATE'),
            ('auditor', 'TEXT'),
            ('ultima_atualizacao', 'DATETIME')
        ]
        
        cursor.execute("PRAGMA table_info(bens)")
        colunas_existentes = [col[1] for col in cursor.fetchall()]
        
        for coluna, tipo in colunas_necessarias:
            if coluna not in colunas_existentes:
                try:
                    cursor.execute(f"ALTER TABLE bens ADD COLUMN {coluna} {tipo}")
                    logger.info(f"Coluna {coluna} adicionada à tabela bens")
                except sqlite3.OperationalError as e:
                    if "duplicate column name" not in str(e).lower():
                        logger.warning(f"Erro ao adicionar coluna {coluna}: {e}")
                    # Coluna já existe, continuar
        
        # Criar índices para performance
        try:
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_numero ON bens(numero)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_situacao ON bens(situacao)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_localizacao ON bens(localizacao)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_bens_nome ON bens(nome)")
        except sqlite3.Error as e:
            logger.warning(f"Erro ao criar índices: {e}")
        
        conn.commit()
        logger.info("Tabela bens verificada/criada com sucesso")
        return True
        
    except Exception as e:
        logger.error(f"Erro ao criar tabela: {str(e)}")
        if conn:
            conn.rollback()
        return False
    finally:
        if conn:
            conn.close()

def buscar_bens(db_path: str, filtros: Dict[str, Any] = None, pagina: int = 1, por_pagina: int = 50) -> Dict[str, Any]:
    """Busca bens com filtros e paginação"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Query base
        query = "SELECT * FROM bens WHERE 1=1"
        params = []
        
        # Aplicar filtros
        if filtros:
            if filtros.get('numero'):
                query += " AND numero LIKE ?"
                params.append(f"%{filtros['numero']}%")
            
            if filtros.get('nome'):
                query += " AND nome LIKE ?"
                params.append(f"%{filtros['nome']}%")
            
            if filtros.get('situacao'):
                query += " AND situacao = ?"
                params.append(filtros['situacao'])
            
            if filtros.get('localizacao'):
                query += " AND localizacao LIKE ?"
                params.append(f"%{filtros['localizacao']}%")
            
            if filtros.get('responsavel'):
                query += " AND responsavel LIKE ?"
                params.append(f"%{filtros['responsavel']}%")
        
        # Ordenação
        query += " ORDER BY numero"
        
        # Contar total
        count_query = f"SELECT COUNT(*) FROM ({query})"
        cursor.execute(count_query, params)
        total = cursor.fetchone()[0]
        
        # Paginação
        offset = (pagina - 1) * por_pagina
        query += " LIMIT ? OFFSET ?"
        params.extend([por_pagina, offset])
        
        # Executar query
        cursor.execute(query, params)
        bens = [dict(row) for row in cursor.fetchall()]
        
        # Converter datas para string
        for bem in bens:
            for key, value in bem.items():
                if isinstance(value, (datetime, date)):
                    bem[key] = value.isoformat()
        
        return {
            'sucesso': True,
            'bens': bens,
            'total': total,
            'pagina': pagina,
            'por_pagina': por_pagina,
            'total_paginas': (total + por_pagina - 1) // por_pagina if por_pagina > 0 else 1
        }
        
    except Exception as e:
        logger.error(f"Erro ao buscar bens: {str(e)}")
        return {
            'sucesso': False,
            'bens': [],
            'total': 0,
            'mensagem': f"Erro ao buscar bens: {str(e)}"
        }
    finally:
        if conn:
            conn.close()

def atualizar_bem(db_path: str, numero_bem: str, dados: Dict[str, Any]) -> Dict[str, Any]:
    """Atualiza os dados de um bem"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Verificar se bem existe
        cursor.execute("SELECT id FROM bens WHERE numero = ?", (numero_bem,))
        if not cursor.fetchone():
            return {
                'sucesso': False,
                'mensagem': f'Bem {numero_bem} não encontrado'
            }
        
        # Construir query dinâmica
        campos = []
        params = []
        
        campos_permitidos = [
            'nome', 'situacao', 'localizacao', 'responsavel',
            'data_ultima_vistoria', 'data_vistoria_atual', 'auditor', 'observacoes'
        ]
        
        for campo in campos_permitidos:
            if campo in dados and dados[campo] is not None:
                campos.append(f"{campo} = ?")
                params.append(dados[campo])
        
        # Se localização foi alterada, atualizar data_localizacao
        if 'localizacao' in dados and dados['localizacao'] is not None:
            campos.append("data_localizacao = CURRENT_TIMESTAMP")
        
        # Sempre atualizar timestamp
        campos.append("ultima_atualizacao = CURRENT_TIMESTAMP")
        
        if not campos:
            return {
                'sucesso': False,
                'mensagem': 'Nenhum campo válido para atualizar'
            }
        
        params.append(numero_bem)
        
        query = f"UPDATE bens SET {', '.join(campos)} WHERE numero = ?"
        cursor.execute(query, params)
        
        conn.commit()
        
        return {
            'sucesso': True,
            'mensagem': f'Bem {numero_bem} atualizado com sucesso'
        }
        
    except sqlite3.Error as e:
        logger.error(f"Erro ao atualizar bem {numero_bem}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro ao atualizar bem: {str(e)}'
        }
    except Exception as e:
        logger.error(f"Erro inesperado ao atualizar bem {numero_bem}: {str(e)}")
        if conn:
            conn.rollback()
        return {
            'sucesso': False,
            'mensagem': f'Erro inesperado: {str(e)}'
        }
    finally:
        if conn:
            conn.close()

def obter_estatisticas(db_path: str) -> Dict[str, Any]:
    """Obtém estatísticas dos bens"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        # Total de bens
        cursor.execute("SELECT COUNT(*) FROM bens")
        total_bens = cursor.fetchone()[0]
        
        # Bens por situação
        cursor.execute("SELECT situacao, COUNT(*) FROM bens GROUP BY situacao")
        situacoes = {row[0]: row[1] for row in cursor.fetchall()}
        
        # Bens por localização (top 10)
        cursor.execute('''
            SELECT localizacao, COUNT(*) 
            FROM bens 
            WHERE localizacao IS NOT NULL AND localizacao != ''
            GROUP BY localizacao 
            ORDER BY COUNT(*) DESC 
            LIMIT 10
        ''')
        localizacoes = {row[0]: row[1] for row in cursor.fetchall()}
        
        # Data da última atualização
        cursor.execute('''
            SELECT MAX(data_criacao) FROM bens
            UNION ALL
            SELECT MAX(ultima_atualizacao) FROM bens WHERE ultima_atualizacao IS NOT NULL
        ''')
        datas = cursor.fetchall()
        datas_validas = [row[0] for row in datas if row[0] is not None]
        ultima_atualizacao = max(datas_validas) if datas_validas else None
        
        return {
            'sucesso': True,
            'estatisticas': {
                'total_bens': total_bens,
                'situacoes': situacoes,
                'localizacoes': localizacoes,
                'ultima_atualizacao': ultima_atualizacao.isoformat() if ultima_atualizacao else None
            }
        }
        
    except Exception as e:
        logger.error(f"Erro ao obter estatísticas: {str(e)}")
        return {
            'sucesso': False,
            'estatisticas': {},
            'mensagem': f"Erro ao obter estatísticas: {str(e)}"
        }
    finally:
        if conn:
            conn.close()

def obter_localizacoes(db_path: str) -> List[str]:
    """Retorna lista de localizações únicas"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT DISTINCT localizacao 
            FROM bens 
            WHERE localizacao IS NOT NULL AND localizacao != ''
            ORDER BY localizacao
        ''')
        
        localizacoes = [row[0] for row in cursor.fetchall()]
        return localizacoes
        
    except Exception as e:
        logger.error(f"Erro ao obter localizações: {str(e)}")
        return []
    finally:
        if conn:
            conn.close()

def obter_situacoes(db_path: str) -> List[str]:
    """Retorna lista de situações únicas"""
    conn = None
    try:
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT DISTINCT situacao 
            FROM bens 
            WHERE situacao IS NOT NULL
            ORDER BY situacao
        ''')
        
        situacoes = [row[0] for row in cursor.fetchall()]
        return situacoes
        
    except Exception as e:
        logger.error(f"Erro ao obter situações: {str(e)}")
        return []
    finally:
        if conn:
            conn.close()

def exportar_para_excel(db_path: str, caminho_saida: str) -> Dict[str, Any]:
    """Exporta dados para arquivo Excel"""
    try:
        import pandas as pd
        
        conn = get_db_connection(db_path)
        
        # Ler dados
        query = '''
            SELECT 
                numero, nome, situacao, localizacao, responsavel,
                data_ultima_vistoria, data_vistoria_atual, auditor,
                observacoes, data_criacao, data_localizacao
            FROM bens 
            ORDER BY numero
        '''
        
        df = pd.read_sql_query(query, conn)
        conn.close()
        
        # Exportar para Excel
        with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Bens', index=False)
            
            # Ajustar largura das colunas
            worksheet = writer.sheets['Bens']
            for idx, col in enumerate(df.columns):
                max_len = max(df[col].astype(str).str.len().max(), len(col)) + 2
                worksheet.column_dimensions[chr(65 + idx)].width = min(max_len, 50)
        
        return {
            'sucesso': True,
            'mensagem': f'Dados exportados para {caminho_saida}',
            'caminho': caminho_saida,
            'total_registros': len(df)
        }
        
    except ImportError:
        return {
            'sucesso': False,
            'mensagem': 'Biblioteca pandas não instalada'
        }
    except Exception as e:
        logger.error(f"Erro ao exportar para Excel: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro ao exportar: {str(e)}'
        }

# ==============================
# FUNÇÃO DE BUSCA RÁPIDA - ADICIONADA
# ==============================

def buscar_bem_rapido(db_path: str, numero_bem: str) -> Dict[str, Any]:
    """
    Busca rápida de bem por número - versão otimizada para verificação
    Retorna apenas dados essenciais
    """
    conn = None
    try:
        if not numero_bem or not numero_bem.strip():
            return {
                'sucesso': False,
                'mensagem': 'Número do bem não informado',
                'encontrado': False
            }
        
        numero_limpo = numero_bem.strip()
        conn = get_db_connection(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT numero, nome, situacao, localizacao
            FROM bens 
            WHERE numero = ?
        ''', (numero_limpo,))
        
        bem = cursor.fetchone()
        
        if bem:
            bem_dict = dict(bem)
            return {
                'sucesso': True,
                'mensagem': 'Bem encontrado',
                'encontrado': True,
                'bem': bem_dict
            }
        else:
            return {
                'sucesso': True,
                'mensagem': f'Bem {numero_limpo} não encontrado',
                'encontrado': False,
                'bem': None
            }
            
    except Exception as e:
        logger.error(f"Erro na busca rápida do bem {numero_bem}: {str(e)}")
        return {
            'sucesso': False,
            'mensagem': f'Erro na busca: {str(e)}',
            'encontrado': False,
            'bem': None
        }
    finally:
        if conn:
            conn.close()

# ==============================
# FIM: db_handler.py COMPLETO FINAL CORRIGIDO
# ==============================